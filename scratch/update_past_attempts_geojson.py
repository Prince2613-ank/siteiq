import json
import openpyxl
import datetime
import shutil
import os

excel_path = r'C:\Users\princ\Downloads\Road 250m parcels (3).xlsx'
geojson_path = r'c:\Users\princ\SiteIQ\apps\web\public\geojson\parcels_250m.geojson'
backup_path = r'c:\Users\princ\SiteIQ\apps\web\public\geojson\parcels_250m.geojson.pre_past_attempts.bak'

# 1. Backup
if not os.path.exists(backup_path):
    shutil.copyfile(geojson_path, backup_path)
    print(f"Created backup at {backup_path}")

wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb.active
headers = [c.value for c in ws[1]]
excel_rows = list(ws.iter_rows(min_row=2, values_only=True))

map_par_idx = headers.index('MAP_PAR_ID')
nh_par_idx = headers.index('NH_PARCEL_ID')
town_idx = headers.index('TOWN')
addr_idx = headers.index('SITE_ADDR')

z_link_idx = headers.index('zillow_link')
z_att_idx = headers.index('zillow_past_attempt')
r_link_idx = headers.index('redfin_link')
r_att_idx = headers.index('redfin_past_attempt')

def normalize_parcel_id(val):
    if val is None:
        return ''
    if isinstance(val, (datetime.date, datetime.datetime)):
        return f"{val.month}-{val.day}"
    return str(val).strip()

excel_by_pid = {}
excel_by_addr_town = {}

for r in excel_rows:
    town = str(r[town_idx]).strip().upper() if r[town_idx] else ''
    addr = str(r[addr_idx]).strip().upper() if r[addr_idx] else ''
    norm_m = normalize_parcel_id(r[map_par_idx])
    norm_nh = normalize_parcel_id(r[nh_par_idx])
    
    if norm_m:
        excel_by_pid[(town, norm_m)] = r
        excel_by_pid[norm_m] = r
    if norm_nh:
        excel_by_pid[(town, norm_nh)] = r
        excel_by_pid[norm_nh] = r
    if town and addr:
        excel_by_addr_town[(town, addr)] = r

with open(geojson_path, 'r', encoding='utf-8') as f:
    gj = json.load(f)

updated_count = 0
z_links_count = 0
r_links_count = 0
z_attempts_count = 0
r_attempts_count = 0

def clean_val(val):
    if val is None:
        return None
    s = str(val).strip()
    if s.lower() in ['none', 'nan', '']:
        return None
    return s

for f in gj['features']:
    props = f['properties']
    p_id = str(props.get('parcel_id', '')).strip()
    town = str(props.get('town', '')).strip().upper()
    addr = str(props.get('address', '')).strip().upper()
    
    row = excel_by_pid.get((town, p_id)) or excel_by_pid.get(p_id) or excel_by_addr_town.get((town, addr))
    if not row:
        print(f"Warning: No match found for parcel {town} {p_id} {addr}")
        continue
    
    updated_count += 1
    
    z_link = clean_val(row[z_link_idx])
    z_att = clean_val(row[z_att_idx])
    r_link = clean_val(row[r_link_idx])
    r_att = clean_val(row[r_att_idx])
    
    props['zillow_link'] = z_link
    props['zillow_past_attempt'] = z_att
    props['redfin_link'] = r_link
    props['redfin_past_attempt'] = r_att
    
    if z_link: z_links_count += 1
    if r_link: r_links_count += 1
    if z_att: z_attempts_count += 1
    if r_att: r_attempts_count += 1

print(f"Total features matched and updated: {updated_count} / {len(gj['features'])}")
print(f"Features with zillow_link: {z_links_count}")
print(f"Features with redfin_link: {r_links_count}")
print(f"Features with zillow_past_attempt: {z_attempts_count}")
print(f"Features with redfin_past_attempt: {r_attempts_count}")

# Save updated GeoJSON
with open(geojson_path, 'w', encoding='utf-8') as f:
    json.dump(gj, f, separators=(',', ':'))

print(f"Successfully saved updated {geojson_path}!")
