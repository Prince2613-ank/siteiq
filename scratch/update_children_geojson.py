import json
import openpyxl
import datetime
import shutil
import os

excel_path = r'C:\Users\princ\Downloads\Road 250m parcels (2).xlsx'
geojson_path = r'c:\Users\princ\SiteIQ\apps\web\public\geojson\parcels_250m.geojson'
backup_path = r'c:\Users\princ\SiteIQ\apps\web\public\geojson\parcels_250m.geojson.pre_children.bak'

# 1. Backup
if not os.path.exists(backup_path):
    shutil.copyfile(geojson_path, backup_path)
    print(f"Created backup at {backup_path}")

wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb.active
headers = [c.value for c in ws[1]]
excel_rows = list(ws.iter_rows(min_row=2, values_only=True))

map_par_idx = headers.index('MAP_PAR_ID')
nh_par_idx = headers.index('NH_PARCEL_ID')
town_idx = headers.index('TOWN')
addr_idx = headers.index('SITE_ADDR')
child_idx = headers.index('children_count_estate_owner')

def normalize_parcel_id(val):
    if val is None:
        return ''
    if isinstance(val, (datetime.date, datetime.datetime)):
        return f"{val.month}-{val.day}"
    return str(val).strip()

excel_by_pid = {}
excel_by_addr_town = {}

for r in excel_rows:
    town = str(r[town_idx]).strip().upper() if r[town_idx] else ''
    addr = str(r[addr_idx]).strip().upper() if r[addr_idx] else ''
    norm_m = normalize_parcel_id(r[map_par_idx])
    norm_nh = normalize_parcel_id(r[nh_par_idx])
    
    if norm_m:
        excel_by_pid[(town, norm_m)] = r
        excel_by_pid[norm_m] = r
    if norm_nh:
        excel_by_pid[(town, norm_nh)] = r
        excel_by_pid[norm_nh] = r
    if town and addr:
        excel_by_addr_town[(town, addr)] = r

with open(geojson_path, 'r', encoding='utf-8') as f:
    gj = json.load(f)

updated_count = 0
assigned_children = 0

for f in gj['features']:
    props = f['properties']
    p_id = str(props.get('parcel_id', '')).strip()
    town = str(props.get('town', '')).strip().upper()
    addr = str(props.get('address', '')).strip().upper()
    
    row = excel_by_pid.get((town, p_id)) or excel_by_pid.get(p_id) or excel_by_addr_town.get((town, addr))
    if not row:
        print(f"Warning: No match found for parcel {town} {p_id} {addr}")
        continue
    
    updated_count += 1
    val = row[child_idx]
    if val is not None and str(val).strip():
        val_clean = str(val).strip()
        props['children_count_estate_owner'] = val_clean
        assigned_children += 1
    else:
        props['children_count_estate_owner'] = None

print(f"Total features matched and updated: {updated_count} / {len(gj['features'])}")
print(f"Features with children_count_estate_owner: {assigned_children}")

# Save updated GeoJSON
with open(geojson_path, 'w', encoding='utf-8') as f:
    json.dump(gj, f, separators=(',', ':'))

print(f"Successfully saved updated {geojson_path}!")
