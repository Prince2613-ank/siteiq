import json
import openpyxl
import datetime
import shutil
import os

excel_path = r'C:\Users\princ\Downloads\Road 250m parcels (1).xlsx'
geojson_path = r'c:\Users\princ\SiteIQ\apps\web\public\geojson\parcels_250m.geojson'
backup_path = r'c:\Users\princ\SiteIQ\apps\web\public\geojson\parcels_250m.geojson.pre_life_event.bak'

# 1. Backup if not already backed up
if not os.path.exists(backup_path):
    shutil.copyfile(geojson_path, backup_path)
    print(f"Created backup at {backup_path}")

wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb.active
headers = [c.value for c in ws[1]]
excel_rows = list(ws.iter_rows(min_row=2, values_only=True))

map_par_idx = headers.index('MAP_PAR_ID')
nh_par_idx = headers.index('NH_PARCEL_ID')
town_idx = headers.index('TOWN')
addr_idx = headers.index('SITE_ADDR')

def normalize_parcel_id(val):
    if val is None:
        return ''
    if isinstance(val, (datetime.date, datetime.datetime)):
        return f"{val.month}-{val.day}"
    return str(val).strip()

excel_by_pid = {}
excel_by_addr_town = {}

for r in excel_rows:
    town = str(r[town_idx]).strip().upper() if r[town_idx] else ''
    addr = str(r[addr_idx]).strip().upper() if r[addr_idx] else ''
    norm_m = normalize_parcel_id(r[map_par_idx])
    norm_nh = normalize_parcel_id(r[nh_par_idx])
    
    if norm_m:
        excel_by_pid[(town, norm_m)] = r
        excel_by_pid[norm_m] = r
    if norm_nh:
        excel_by_pid[(town, norm_nh)] = r
        excel_by_pid[norm_nh] = r
    if town and addr:
        excel_by_addr_town[(town, addr)] = r

with open(geojson_path, 'r', encoding='utf-8') as f:
    gj = json.load(f)

# Columns to extract from Excel into GeoJSON properties
cols_to_map = [
    'owner_name',
    'owner_type',
    'owner_mailing_address',
    'owner_lives_on_lot',
    'owner_distance_miles',
    'owner_distance_group',
    'years_owned',
    'last_sale_date',
    'last_sale_price',
    'deed_book_page',
    'town_record_link',
    'company_people',
    'company_record_link',
    'tax_takings_found',
    'last_tax_taking',
    'unpaid_tax_signal',
    'tax_check_link',
    'site_address_town_record',
    'life_event_signal',
    'estate_record_detail',
    'past_listing_attempt'
]

col_indices = {c: headers.index(c) for c in cols_to_map if c in headers}

updated_count = 0
life_event_assigned = 0
estate_detail_assigned = 0

for f in gj['features']:
    props = f['properties']
    p_id = str(props.get('parcel_id', '')).strip()
    town = str(props.get('town', '')).strip().upper()
    addr = str(props.get('address', '')).strip().upper()
    
    row = excel_by_pid.get((town, p_id)) or excel_by_pid.get(p_id) or excel_by_addr_town.get((town, addr))
    if not row:
        print(f"Warning: No match found for parcel {town} {p_id} {addr}")
        continue
    
    updated_count += 1
    for col_name, idx in col_indices.items():
        val = row[idx]
        if val is None:
            # If property already exists with value and val is None, don't overwrite if it was already set, or set to null
            if col_name in ['life_event_signal', 'estate_record_detail']:
                props[col_name] = None
        else:
            if isinstance(val, (datetime.date, datetime.datetime)):
                val_str = val.strftime('%b %d, %Y')
            elif isinstance(val, float) and val.is_integer() and col_name in ['tax_takings_found']:
                val_str = int(val)
            elif isinstance(val, str):
                val_str = val.strip()
            else:
                val_str = val
            props[col_name] = val_str
            
            if col_name == 'life_event_signal' and val_str:
                life_event_assigned += 1
            if col_name == 'estate_record_detail' and val_str:
                estate_detail_assigned += 1

print(f"Total features updated: {updated_count} / {len(gj['features'])}")
print(f"Features with life_event_signal: {life_event_assigned}")
print(f"Features with estate_record_detail: {estate_detail_assigned}")

# Save updated GeoJSON
with open(geojson_path, 'w', encoding='utf-8') as f:
    json.dump(gj, f, separators=(',', ':'))

print(f"Successfully saved updated {geojson_path}!")
